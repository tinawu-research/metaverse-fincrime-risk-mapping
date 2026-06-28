from pathlib import Path

from docx import Document
from docx.shared import Pt


SOURCE = Path(
    r"C:\Users\mtiwar05\OneDrive - Charles Sturt University\CSU\Metaverse and Cryptocurrencies\Analysis\Codex\Results\Round 5\mirofish_local_qwen_free_analysis_20260527-112059\enhanced_complete_analysis_report.md"
)
TARGET = SOURCE.with_suffix(".docx")


def clean_inline(text):
    return text.replace("**", "").replace("`", "")


doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line:
        i += 1
        continue

    if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].startswith("|"):
            table_lines.append(lines[i])
            i += 1
        rows = []
        for idx, table_line in enumerate(table_lines):
            cells = [clean_inline(cell.strip()) for cell in table_line.strip("|").split("|")]
            if idx == 1 and all(set(cell.replace(":", "").replace("-", "").strip()) == set() for cell in cells):
                continue
            rows.append(cells)
        if rows:
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
        continue

    if line.startswith("# "):
        doc.add_heading(clean_inline(line[2:]), level=1)
    elif line.startswith("## "):
        doc.add_heading(clean_inline(line[3:]), level=2)
    elif line.startswith("### "):
        doc.add_heading(clean_inline(line[4:]), level=3)
    elif line.startswith("- "):
        doc.add_paragraph(clean_inline(line[2:]), style="List Bullet")
    elif line and line[0].isdigit() and ". " in line[:5]:
        doc.add_paragraph(clean_inline(line.split(". ", 1)[1]), style="List Number")
    else:
        doc.add_paragraph(clean_inline(line))
    i += 1

doc.save(TARGET)
print(TARGET)
