from pathlib import Path

from docx import Document
from docx.shared import Pt


OUT_DIR = Path(
    r"C:\Users\mtiwar05\OneDrive - Charles Sturt University\CSU\Metaverse and Cryptocurrencies\Analysis\Codex\Results\Round 5\mirofish_local_qwen_free_analysis_20260527-112059"
)

md_path = OUT_DIR / "topic_model_emerging_threats_modelling_intervention_focus.md"
docx_path = OUT_DIR / "topic_model_emerging_threats_modelling_intervention_focus.docx"
svg_path = OUT_DIR / "topic_model_threat_intervention_framework.svg"
mmd_path = OUT_DIR / "topic_model_threat_intervention_framework.mmd"

report = """# Topic-Modelling Focus: Emerging Threats, Modelling, Analysis, and Intervention Points

## Reframed Contribution

The Round 5 topic modelling should be interpreted as more than a thematic map of metaverse and cryptocurrency scholarship. It identifies **emerging threat environments** at the intersection of metaverse platforms and cryptocurrency ecosystems. The strongest contribution of the analysis is therefore not simply that metaverse and crypto features create opportunities for financial crime, but that the topic structure helps organise how those opportunities can be **modelled, analysed, and converted into intervention points**.

This reframing keeps the analysis aligned with the research question:

> How do the combined features of metaverse platforms and cryptocurrency ecosystems generate new money laundering and other financial crime opportunities, and subsequently help model, analyse, and identify intervention points within these emerging financial crime environments?

Under this framing, the topic model performs three analytical functions. First, it identifies the main domains in which emerging threats cluster. Second, it shows which domains are connected across technical, financial, governance, identity, and security layers. Third, it provides a structure for deciding where interventions should be placed.

## Topic-Modelling Basis

The Round 5 LDA model retained **seven topics** using a coherence-first decision rule. The corpus comprised **59 PDFs**, **545,779 extracted words**, **247,995 cleaned tokens**, and **332 modelling/scoring chunks**. The selected model was **K = 7**, with c_v coherence of **0.418574**, perplexity of **275.029468**, and log-likelihood bound of **-1,596,070.22**.

The seven topics are best read as a layered map of emerging financial crime environments:

| Topic | Topic label | Prevalence | Dominant PDFs | Threat-environment interpretation |
|---|---:|---:|---:|---|
| T6 | Crypto finance, DeFi, scams, and grey-economy risks | 19.17% | 13 | Core money-laundering, scam, grey-economy, and illicit-finance environment. |
| T2 | Avatar privacy, cyberattack detection, and metaverse security | 18.53% | 10 | Identity, avatar, behavioural, cyberattack, and security-monitoring environment. |
| T1 | Blockchain-based metaverse governance, ownership, and token infrastructure | 16.30% | 8 | Governance, ownership, token infrastructure, trust, and protocol environment. |
| T7 | Cryptographic access control, authentication, and encryption protocols | 13.64% | 9 | Authentication, encryption, access-control, transfer-transaction, and privacy-preserving-control environment. |
| T5 | Blockchain gaming, NFT markets, and digital-property law | 13.37% | 7 | NFT market, gaming economy, jurisdiction, property, volatility, and Sybil-wallet environment. |
| T3 | Cryptocurrency networks, wallets, exchanges, and regulatory risk | 9.65% | 7 | Wallet, exchange, Bitcoin-network, anonymity, illicit-finance, Tornado Cash, and regulatory-risk environment. |
| T4 | Digital payments, privacy policy, trust, and fraud detection | 9.34% | 5 | Payment, privacy-policy, off-chain transaction, fraud-detection, gas-fee, and Sybil-attack environment. |

## Emerging Threat Environments Identified by the Topic Model

### 1. Crypto Finance, DeFi, Scams, and Grey-Economy Risks

T6 is the largest topic and is the clearest direct anchor for money laundering and financial crime. It identifies the threat environment in which cryptocurrency, digital currency, DeFi, scams, social media, and grey-economy activity intersect. This topic indicates that metaverse financial crime should not be modelled only as a technical cybersecurity issue; it also involves fraud scripts, social influence, speculative markets, informal economies, and offender adaptation.

For modelling, this topic points to scam typology construction, wallet-flow analysis, fraud network mapping, and social-signal monitoring. For intervention, it points to scam warnings, suspicious wallet clustering, off-ramp due diligence, platform moderation, and cross-platform reporting of fraudulent campaigns.

### 2. Avatar Privacy, Cyberattack Detection, and Metaverse Security

T2 identifies a second threat environment around avatars, attacks, prediction, detection, cybersecurity, authentication, immersive behaviour, nodes, and behavioural patterns. This topic shifts the analysis from crypto transactions alone to identity-mediated financial crime. In metaverse settings, avatars are not just visual representations; they can become trust-bearing, asset-holding, social-engineering, and authentication targets.

For modelling, T2 supports avatar-wallet-account linkage, behavioural anomaly detection, cyberattack prediction, credential-compromise analysis, and identity-risk scoring. For intervention, it points to risk-based re-authentication, phishing detection, wallet-drain alerts, avatar impersonation controls, behavioural monitoring, and identity-recovery procedures.

### 3. Blockchain Governance, Ownership, and Token Infrastructure

T1 captures the governance and infrastructure layer: trust, governance, ownership, blockchain networks, tokens, swaps, infrastructure, and protocols. This threat environment matters because financial crime risk often emerges from control points embedded in protocols, DAOs, token issuance, asset ownership systems, and cross-platform infrastructure.

For modelling, T1 supports governance-risk graphs, smart-contract dependency mapping, ownership-chain analysis, DAO treasury monitoring, and token-infrastructure risk scoring. For intervention, it points to governance transparency, smart-contract audits, multisig controls, project disclosure, treasury-monitoring triggers, and accountability controls for token issuers and virtual organisations.

### 4. Cryptographic Access Control, Authentication, and Encryption Protocols

T7 shows that intervention cannot rely only on post-transaction detection. Access control, authentication protocols, cryptographic key management, privacy-preserving technology, and on-chain transfer mechanisms form a preventive control environment. The threat is not only misuse of assets after acquisition, but also compromise of access rights, keys, credentials, and transfer permissions.

For modelling, T7 supports authentication-failure analysis, key-management risk mapping, access-control policy modelling, credential-compromise pathways, and secure-transfer process modelling. For intervention, it points to multi-factor authentication, hardware or custodial safeguards for high-value assets, smart-contract permission review, recovery protocols, and privacy-preserving compliance mechanisms.

### 5. Blockchain Gaming, NFT Markets, and Digital-Property Law

T5 maps a threat environment around gaming, NFTs, jurisdiction, property, intellectual property, volatility, and Sybil wallets. This is where criminal opportunity is linked to asset valuation, market opacity, digital-property rights, jurisdictional uncertainty, and gaming economies. NFT and virtual-property markets create plausible mechanisms for overvaluation, wash trading, sham sale, account farming, and illicit integration.

For modelling, T5 supports NFT wash-trading detection, virtual-property valuation anomaly analysis, Sybil-wallet identification, marketplace manipulation metrics, and jurisdictional-risk mapping. For intervention, it points to marketplace surveillance, provenance standards, high-value transaction checks, related-wallet detection, intellectual-property review, and stronger reporting protocols for suspicious virtual-asset trades.

### 6. Cryptocurrency Networks, Wallets, Exchanges, and Regulatory Risk

T3 identifies the regulated and semi-regulated crypto-financial interface: wallets, exchanges, Bitcoin networks, anonymity/privacy, illicit finance, consumer protection, addresses, and Tornado Cash. This topic is critical because many metaverse-crypto crimes eventually require liquidity, exchange access, stablecoin movement, or conversion into fiat.

For modelling, T3 supports wallet-cluster analysis, exchange interaction mapping, exposure to high-risk services, mixer/bridge/pathway tracing, and regulatory-risk scoring. For intervention, it points to exchange due diligence, wallet screening, sanctions and mixer exposure alerts, Travel Rule-style data sharing where applicable, and off-ramp review.

### 7. Digital Payments, Privacy Policy, Trust, and Fraud Detection

T4 isolates the payment and fraud-detection layer: privacy policy, digital payments, fraud detection, gas fees, data flows, off-chain transactions, trust management, and Sybil attack. This topic shows that intervention must consider both on-chain and off-chain environments. Fraud may be executed through off-chain communications and platform-level trust cues, while value transfer occurs on-chain or through connected payment systems.

For modelling, T4 supports fraud-detection models that combine transaction metadata, platform logs, privacy-policy exposure, payment-flow data, and Sybil indicators. For intervention, it points to payment monitoring, transparent privacy policies, platform-level fraud reporting, off-chain/on-chain data linkage, and Sybil-resistance mechanisms.

## Modelling Framework

The topic model supports a **multi-layer threat-modelling framework**. Each LDA topic can be treated as a layer or module in an emerging financial crime environment.

| Model layer | Topic anchor | Main entities | Main relationships | Analytical purpose |
|---|---|---|---|---|
| Social-deception layer | T6, T2 | avatars, victims, promoters, communities, social media accounts | influences, impersonates, recruits, deceives | Identify scams, grooming, trust manipulation, and victim pathways. |
| Identity-access layer | T2, T7 | avatars, wallets, keys, credentials, devices, authentication events | controls, compromises, authenticates, transfers | Detect account takeover, identity theft, wallet compromise, and credential abuse. |
| Asset-market layer | T5, T1 | NFTs, virtual land, game assets, tokens, smart contracts, marketplaces | mints, lists, buys, sells, swaps, owns | Detect wash trading, overvaluation, sham sale, suspicious asset cycling, and market manipulation. |
| Transaction-flow layer | T3, T4, T6 | wallets, exchanges, bridges, mixers, DeFi pools, stablecoins | sends, receives, swaps, bridges, off-ramps | Trace laundering, layering, chain-hopping, mixer exposure, and fiat conversion. |
| Governance-control layer | T1, T7 | DAOs, protocols, multisig wallets, platform operators, regulators | governs, audits, approves, enforces | Identify weak governance, treasury risk, protocol abuse, and points of accountability. |
| Detection-intervention layer | T2, T4, T7 | risk models, alerts, reports, compliance controls, law-enforcement actions | flags, blocks, escalates, investigates | Convert signals into prevention, monitoring, reporting, and enforcement actions. |

This framework helps avoid a narrow reading of financial crime as only transaction movement. In metaverse-cryptocurrency environments, offending can begin in social interaction, move through identity compromise or asset promotion, become encoded in virtual asset transactions, pass through crypto infrastructure, and only later appear at an exchange or fiat off-ramp.

## Analytical Indicators by Threat Environment

| Threat environment | Indicators for analysis |
|---|---|
| Scam and grey-economy finance | repeated promotional activity, high-return claims, victim reports, wallet clusters linked to scam campaigns, rapid movement to exchanges or bridges. |
| Avatar and identity abuse | sudden device or login changes, high-value asset transfers after credential reset, avatar impersonation, wallet-drain patterns, phishing-link exposure. |
| Governance and token infrastructure | concentrated DAO voting, treasury transfers to new wallets, unaudited smart contracts, opaque project control, suspicious token allocation. |
| Authentication and access control | failed authentication bursts, key compromise events, abnormal permission grants, risky smart-contract approvals, unusual transfer permissions. |
| NFT markets and virtual property | self-funded trades, circular NFT movement, abnormal price spikes, related-wallet counterparties, suspicious metadata/provenance, Sybil wallet behaviour. |
| Wallets, exchanges, and regulatory risk | mixer or Tornado Cash exposure, bridge chains, rapid stablecoin conversion, high-risk jurisdiction signals, exchange off-ramp attempts after suspicious NFT/DeFi activity. |
| Digital payments and fraud detection | off-chain payment promises, privacy-policy gaps, Sybil attack patterns, gas-fee anomalies, suspicious off-chain/on-chain timing relationships. |

## Intervention Points

The intervention model should follow the lifecycle of emerging metaverse-crypto financial crime:

1. **Platform and protocol design:** require auditability, suspicious activity logging, secure smart-contract design, provenance standards, and privacy-preserving compliance controls.
2. **Onboarding and identity:** apply risk-based verification to high-value creators, sellers, financial-service operators, DAO controllers, and repeat high-risk traders.
3. **Asset issuance and listing:** screen NFT collections, virtual land, token launches, and digital twins for provenance, suspicious valuation, cloned content, and hidden control.
4. **Market and transaction activity:** detect wash trading, circular wallet flows, bridge layering, mixer exposure, Sybil wallets, rapid resale, abnormal pricing, and asset cycling.
5. **Social and behavioural monitoring:** monitor scam campaigns, fake events, impersonation, account takeover, phishing, and influencer-style financial promotions.
6. **Cross-chain and DeFi movement:** apply risk scoring to bridges, DEXs, liquidity pools, stablecoin conversions, and high-risk services.
7. **Exchange and fiat off-ramp:** integrate metaverse-specific risk indicators into KYC, transaction monitoring, enhanced due diligence, and suspicious matter reporting.
8. **Post-event investigation:** preserve wallet histories, asset metadata, marketplace logs, avatar-account records, chat evidence, smart-contract addresses, and off-ramp data.

## Revised Manuscript Argument

The topic modelling indicates that emerging financial crime risks at the metaverse-cryptocurrency intersection are not confined to a single technological domain. Rather, they form a multi-layered threat environment in which crypto-finance and grey-economy risks (T6), avatar privacy and cyberattack detection (T2), blockchain governance and token infrastructure (T1), cryptographic access control (T7), NFT markets and digital-property law (T5), wallet/exchange regulatory risk (T3), and digital-payment fraud detection (T4) operate as interdependent layers.

This layered topic structure helps move the analysis from description to modelling. Each topic identifies a class of entities, relationships, behavioural indicators, and control points. Together, the topics support a graph-based and lifecycle-based model of emerging financial crime environments. This model links avatars, wallets, NFTs, virtual land, smart contracts, exchanges, DAOs, bridges, platform accounts, social interactions, and regulatory controls. It also helps identify intervention points before, during, and after suspicious activity: secure-by-design architecture, risk-based onboarding, asset-listing controls, transaction monitoring, behavioural anomaly detection, cross-chain analytics, off-ramp due diligence, and evidence-preserving investigation.

The main contribution is therefore that topic modelling identifies the structure of emerging threats and provides a basis for modelling how money laundering and related financial crimes can be detected and disrupted in metaverse-cryptocurrency ecosystems.
"""

mermaid = """flowchart LR
    T6["T6 Crypto finance, DeFi, scams, grey-economy risks"] --> M["Emerging metaverse-crypto financial crime environment"]
    T2["T2 Avatar privacy, cyberattack detection, security"] --> M
    T1["T1 Governance, ownership, token infrastructure"] --> M
    T7["T7 Access control, authentication, encryption"] --> M
    T5["T5 Gaming, NFT markets, digital-property law"] --> M
    T3["T3 Wallets, exchanges, regulatory risk"] --> M
    T4["T4 Payments, privacy policy, fraud detection"] --> M
    M --> A["Model: multi-layer graph of avatars, wallets, assets, contracts, platforms, exchanges"]
    A --> B["Analyse: typologies, anomaly signals, entity links, transaction paths, behavioural patterns"]
    B --> C["Intervene: design controls, onboarding, marketplace review, transaction monitoring, cross-chain analytics, off-ramp due diligence, investigation"]
"""

svg = """<svg xmlns="http://www.w3.org/2000/svg" width="1400" height="880" viewBox="0 0 1400 880">
  <style>
    .topic { fill:#f7f7f7; stroke:#333; stroke-width:1.4; rx:8; }
    .core { fill:#fff3cd; stroke:#9a6b00; stroke-width:2; rx:8; }
    .stage { fill:#e9f5ff; stroke:#185a8d; stroke-width:1.8; rx:8; }
    .label { font-family:Arial, sans-serif; font-size:18px; fill:#111; }
    .small { font-family:Arial, sans-serif; font-size:15px; fill:#222; }
    .title { font-family:Arial, sans-serif; font-size:26px; font-weight:bold; fill:#111; }
    .arrow { stroke:#333; stroke-width:1.6; marker-end:url(#arrow); }
  </style>
  <defs>
    <marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto">
      <path d="M0,0 L0,6 L9,3 z" fill="#333"/>
    </marker>
  </defs>
  <text x="70" y="50" class="title">Topic Model to Threat Modelling and Intervention Framework</text>

  <rect x="60" y="90" width="370" height="70" class="topic"/><text x="78" y="120" class="label">T6 Crypto finance, DeFi, scams</text><text x="78" y="145" class="small">Grey economy, ML/TF, scam pathways</text>
  <rect x="60" y="180" width="370" height="70" class="topic"/><text x="78" y="210" class="label">T2 Avatar privacy and cyberattack</text><text x="78" y="235" class="small">Identity, behaviour, detection, security</text>
  <rect x="60" y="270" width="370" height="70" class="topic"/><text x="78" y="300" class="label">T1 Governance and token infrastructure</text><text x="78" y="325" class="small">Ownership, protocols, trust, swaps</text>
  <rect x="60" y="360" width="370" height="70" class="topic"/><text x="78" y="390" class="label">T7 Access control and authentication</text><text x="78" y="415" class="small">Keys, encryption, permissions, on-chain transfer</text>
  <rect x="60" y="450" width="370" height="70" class="topic"/><text x="78" y="480" class="label">T5 Gaming, NFTs, digital property</text><text x="78" y="505" class="small">Markets, jurisdiction, volatility, Sybil wallets</text>
  <rect x="60" y="540" width="370" height="70" class="topic"/><text x="78" y="570" class="label">T3 Wallets, exchanges, regulation</text><text x="78" y="595" class="small">Anonymity, illicit finance, off-ramp risk</text>
  <rect x="60" y="630" width="370" height="70" class="topic"/><text x="78" y="660" class="label">T4 Payments, privacy, fraud detection</text><text x="78" y="685" class="small">Off-chain transactions, trust, Sybil attacks</text>

  <rect x="520" y="300" width="360" height="150" class="core"/>
  <text x="545" y="335" class="label">Emerging metaverse-crypto</text>
  <text x="545" y="362" class="label">financial crime environment</text>
  <text x="545" y="395" class="small">A layered socio-technical setting where</text>
  <text x="545" y="417" class="small">identity, assets, transactions, governance,</text>
  <text x="545" y="439" class="small">and controls interact.</text>

  <rect x="990" y="170" width="320" height="100" class="stage"/>
  <text x="1015" y="205" class="label">Model</text>
  <text x="1015" y="233" class="small">Graph avatars, wallets, assets,</text>
  <text x="1015" y="254" class="small">contracts, platforms, exchanges</text>
  <rect x="990" y="340" width="320" height="100" class="stage"/>
  <text x="1015" y="375" class="label">Analyse</text>
  <text x="1015" y="403" class="small">Typologies, anomalies, entity links,</text>
  <text x="1015" y="424" class="small">transaction paths, behaviour</text>
  <rect x="990" y="510" width="320" height="130" class="stage"/>
  <text x="1015" y="545" class="label">Intervene</text>
  <text x="1015" y="573" class="small">Design controls, onboarding, marketplace</text>
  <text x="1015" y="594" class="small">review, monitoring, cross-chain analytics,</text>
  <text x="1015" y="615" class="small">off-ramp due diligence, investigation</text>

  <line x1="430" y1="125" x2="520" y2="350" class="arrow"/>
  <line x1="430" y1="215" x2="520" y2="360" class="arrow"/>
  <line x1="430" y1="305" x2="520" y2="370" class="arrow"/>
  <line x1="430" y1="395" x2="520" y2="380" class="arrow"/>
  <line x1="430" y1="485" x2="520" y2="390" class="arrow"/>
  <line x1="430" y1="575" x2="520" y2="400" class="arrow"/>
  <line x1="430" y1="665" x2="520" y2="410" class="arrow"/>
  <line x1="880" y1="350" x2="990" y2="220" class="arrow"/>
  <line x1="880" y1="375" x2="990" y2="390" class="arrow"/>
  <line x1="880" y1="410" x2="990" y2="575" class="arrow"/>
</svg>
"""

md_path.write_text(report, encoding="utf-8")
mmd_path.write_text(mermaid, encoding="utf-8")
svg_path.write_text(svg, encoding="utf-8")

doc = Document()
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)

for line in report.splitlines():
    if not line.strip():
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:], level=1)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=2)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=3)
    elif line.startswith("|"):
        # Tables are retained in the markdown version; keep Word simple/readable.
        doc.add_paragraph(line)
    elif line.startswith("- "):
        doc.add_paragraph(line[2:], style="List Bullet")
    elif line[:2].isdigit() and ". " in line[:5]:
        doc.add_paragraph(line.split(". ", 1)[1], style="List Number")
    else:
        doc.add_paragraph(line.replace("**", "").replace("`", ""))

doc.save(docx_path)
print(md_path)
print(docx_path)
print(svg_path)
print(mmd_path)
